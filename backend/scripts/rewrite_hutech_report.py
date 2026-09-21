from __future__ import annotations

import argparse
import json
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "bao_cao_do_an_co_so_HUTECH.docx"

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(13)
H1_SIZE = Pt(16)
H2_SIZE = Pt(14)
H3_SIZE = Pt(14)
ACCENT = RGBColor(31, 78, 121)
LIGHT_FILL = "EAF2F8"
TABLE_HEADER_FILL = "D9EAF7"


TOC_ENTRIES = [
    ("LỜI CAM ĐOAN", 0),
    ("DANH MỤC CÁC KÝ HIỆU, CHỮ VIẾT TẮT", 0),
    ("DANH MỤC CÁC BẢNG", 0),
    ("DANH MỤC CÁC HÌNH VẼ, ĐỒ THỊ", 0),
    ("Chương 1. TỔNG QUAN", 0),
    ("1.1 Nhiệm vụ đồ án", 1),
    ("1.2 Cấu trúc đồ án", 1),
    ("1.3 Các nghiên cứu liên quan", 1),
    ("1.4 Yêu cầu đặt ra cho hệ thống", 1),
    ("Chương 2. CƠ SỞ LÝ THUYẾT", 0),
    ("2.1 Đặc điểm tiếng Việt trong văn bản chat", 1),
    ("2.2 Cơ sở xử lý và chuẩn hóa văn bản", 1),
    ("2.3 Phương pháp khôi phục dấu tiếng Việt", 1),
    ("2.4 Phương pháp giải quyết đa nghĩa", 1),
    ("2.5 Kiến trúc hệ thống phần mềm", 1),
    ("2.6 Công nghệ sử dụng và ràng buộc triển khai", 1),
    ("2.7 Cơ sở đánh giá kết quả", 1),
    ("Chương 3. KẾT QUẢ THỰC NGHIỆM", 0),
    ("3.1 Thiết kế hệ thống", 1),
    ("3.2 Cài đặt pipeline chuẩn hóa", 1),
    ("3.3 Cài đặt dữ liệu và tài nguyên khôi phục dấu", 1),
    ("3.4 Cài đặt giao diện người dùng", 1),
    ("3.5 Kết quả kiểm thử và đánh giá", 1),
    ("3.6 Nhận xét kết quả thực nghiệm", 1),
    ("Chương 4. KẾT LUẬN VÀ KIẾN NGHỊ", 0),
    ("4.1 Kết luận", 1),
    ("4.2 Kiến nghị", 1),
    ("TÀI LIỆU THAM KHẢO", 0),
    ("PHỤ LỤC", 0),
]


TABLE_LIST = [
    "Bảng 2.1 So sánh các phương pháp chuẩn hóa văn bản tiếng Việt",
    "Bảng 2.2 Các chỉ số đánh giá được dùng trong đồ án",
    "Bảng 3.1 Công nghệ và thư viện sử dụng trong dự án",
    "Bảng 3.2 Các bảng cơ sở dữ liệu chính trong hệ thống",
    "Bảng 3.3 Các endpoint REST API chính",
    "Bảng 3.4 Kết quả đánh giá khôi phục dấu trên tập kiểm tra",
    "Bảng 3.5 Kết quả regression trên tập câu chat thực tế",
    "Bảng 3.6 Tổng hợp kết quả kiểm thử phần mềm",
    "Bảng 3.7 Cách đối chiếu các chỉ số phần mềm",
    "Bảng 3.8 Đối chiếu yêu cầu ban đầu và kết quả thực hiện",
    "Bảng B.1 Ví dụ chuẩn hóa tiêu biểu",
]


FIGURE_LIST = [
    "Hình 3.1 Kiến trúc tổng thể hệ thống",
    "Hình 3.2 Pipeline chuẩn hóa văn bản",
    "Hình 3.3 Luồng xử lý yêu cầu chuẩn hóa theo thời gian gần thực",
    "Hình 3.4 Sơ đồ quan hệ dữ liệu rút gọn",
    "Hình 3.5 Giao diện Workspace của người dùng",
    "Hình 3.6 Giao diện quản lý từ điển và kiểm duyệt",
]


def set_run_font(run, size: Pt | None = None, bold: bool | None = None, italic: bool | None = None):
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name in ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.5

    h1 = styles["Heading 1"]
    h1.font.size = H1_SIZE
    h1.font.bold = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)

    h2 = styles["Heading 2"]
    h2.font.size = H2_SIZE
    h2.font.bold = True
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)

    h3 = styles["Heading 3"]
    h3.font.size = H3_SIZE
    h3.font.italic = True
    h3.font.bold = False
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(6)

    h4 = styles["Heading 4"]
    h4.font.size = BODY_SIZE
    h4.font.underline = True
    h4.font.bold = False
    h4.paragraph_format.space_before = Pt(6)
    h4.paragraph_format.space_after = Pt(4)


def add_page_number(section):
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    set_run_font(run, Pt(11))
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)


def restart_page_number(section, start: int = 1):
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def add_paragraph(
    doc: Document,
    text: str = "",
    style: str | None = None,
    align: WD_ALIGN_PARAGRAPH | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    size: Pt | None = None,
    first_line: bool = True,
    space_after: int = 6,
):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    elif style in {"Heading 1"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(space_after)
    if first_line and not style:
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size or BODY_SIZE, bold=bold, italic=italic)
    return paragraph


def add_heading(doc: Document, text: str, level: int, page_break_before: bool = True):
    style = f"Heading {level}"
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.5
    if level == 1 and page_break_before:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    if level == 1:
        set_run_font(run, H1_SIZE, bold=True)
    elif level == 2:
        set_run_font(run, H2_SIZE, bold=True)
    elif level == 3:
        set_run_font(run, H3_SIZE, italic=True)
    else:
        set_run_font(run, BODY_SIZE)
        run.font.underline = True
    return paragraph


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float] | None = None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[cell_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.2
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_font(run, Pt(12), bold=(row_index == 0))
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)


def style_index_table(table, widths: list[float] | None = None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=20, start=60, bottom=20, end=60)
            if widths:
                cell.width = Inches(widths[cell_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, Pt(11.5))


def add_caption(doc: Document, text: str, above: bool = False):
    paragraph = add_paragraph(
        doc,
        text,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        first_line=False,
        space_after=8 if above else 10,
    )
    paragraph.paragraph_format.space_before = Pt(6 if not above else 4)
    return paragraph


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]):
    if caption:
        add_caption(doc, caption, above=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            row.cells[index].text = value
    style_table(table, widths)
    add_paragraph(doc, "", first_line=False, space_after=6)
    return table


def font_path():
    for candidate in [
        "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]:
        if Path(candidate).exists():
            return candidate
    return None


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, text, font, fill="#F7FBFF", outline="#2F5597"):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    lines = wrap_text(draw, text, font, x2 - x1 - 28)
    total_h = len(lines) * 26
    y = y1 + ((y2 - y1 - total_h) // 2)
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, fill="#14213D", font=font)
        y += 26


def arrow(draw, start, end, color="#486581"):
    draw.line([start, end], fill=color, width=4)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - sign * 14, ey - 8), (ex - sign * 14, ey + 8)]
    else:
        sign = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 8, ey - sign * 14), (ex + 8, ey - sign * 14)]
    draw.polygon(points, fill=color)


def diagram_image(kind: str) -> BytesIO:
    path = font_path()
    font = ImageFont.truetype(path, 23) if path else ImageFont.load_default()
    small = ImageFont.truetype(path, 19) if path else ImageFont.load_default()
    title_font = ImageFont.truetype(path, 28) if path else ImageFont.load_default()
    img = Image.new("RGB", (1280, 720), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 28), kind, fill="#1F4E79", font=title_font)

    if kind == "Kiến trúc tổng thể hệ thống":
        boxes = [
            ((55, 210, 265, 345), "Người dùng\ntrình duyệt"),
            ((360, 190, 590, 365), "React và Vite\nphần giao diện"),
            ((700, 190, 930, 365), "FastAPI\nphần xử lý"),
            ((1030, 95, 1225, 245), "SQL Server\nlưu dữ liệu"),
            ((1030, 315, 1225, 465), "Gemini\nkiểm tra khi cần"),
        ]
        for xy, text in boxes:
            draw_box(draw, xy, text.replace("\n", " "), font)
        arrow(draw, (265, 278), (360, 278))
        arrow(draw, (590, 278), (700, 278))
        arrow(draw, (930, 245), (1030, 170))
        arrow(draw, (930, 318), (1030, 390))
        draw.text((380, 410), "Kết quả trả về theo JSON, giao diện hiển thị gần như ngay khi người dùng gõ.", fill="#334E68", font=small)
    elif kind == "Pipeline chuẩn hóa văn bản":
        labels = [
            "Văn bản nhập",
            "Chuẩn Unicode",
            "Bỏ emoji",
            "Khôi phục dấu",
            "Mở rộng viết tắt",
            "Kiểm tra nghĩa",
            "Kết quả",
        ]
        x = 55
        for label in labels:
            draw_box(draw, (x, 245, x + 150, 355), label, small)
            if label != labels[-1]:
                arrow(draw, (x + 150, 300), (x + 190, 300))
            x += 190
        draw.text((80, 420), "Các bước đơn giản được tách thành từng lớp để dễ kiểm thử, thay đổi và mở rộng từ điển.", fill="#334E68", font=small)
    elif kind == "Luồng xử lý yêu cầu chuẩn hóa theo thời gian gần thực":
        lanes = [("Người dùng", 90), ("Frontend", 330), ("Backend", 570), ("Dữ liệu", 810), ("AI khi cần", 1050)]
        for label, x in lanes:
            draw.text((x - 55, 95), label, fill="#1F4E79", font=font)
            draw.line((x, 135, x, 620), fill="#CBD5E1", width=3)
        steps = [
            ((90, 180), (330, 180), "nhập văn bản"),
            ((330, 250), (570, 250), "gửi yêu cầu"),
            ((570, 320), (810, 320), "đọc từ điển"),
            ((570, 390), (1050, 390), "kiểm tra câu khó"),
            ((570, 470), (330, 470), "trả kết quả"),
            ((330, 545), (90, 545), "hiển thị"),
        ]
        for start, end, text in steps:
            arrow(draw, start, end)
            x = min(start[0], end[0]) + 30
            y = start[1] - 28
            draw.text((x, y), text, fill="#334E68", font=small)
    elif kind == "Sơ đồ quan hệ dữ liệu rút gọn":
        boxes = [
            ((70, 150, 300, 275), "users\nthông tin tài khoản"),
            ((390, 150, 620, 275), "user_sessions\nphiên đăng nhập"),
            ((710, 150, 940, 275), "normalization_history\nlịch sử xử lý"),
            ((70, 390, 300, 515), "abbreviations\ntừ viết tắt"),
            ((390, 390, 620, 515), "pending_abbreviations\nchờ duyệt"),
            ((710, 390, 940, 515), "phrase_overrides\ncụm từ ưu tiên"),
            ((1010, 270, 1230, 395), "usage_stats\nthống kê sử dụng"),
        ]
        for xy, text in boxes:
            draw_box(draw, xy, text.replace("\n", " "), small)
        arrow(draw, (300, 210), (390, 210))
        arrow(draw, (300, 450), (390, 450))
        arrow(draw, (620, 450), (710, 450))
        arrow(draw, (940, 330), (1010, 330))
        arrow(draw, (300, 210), (710, 210))
    elif kind == "Giao diện Workspace của người dùng":
        draw.rounded_rectangle((70, 120, 1210, 615), radius=22, outline="#2F5597", width=4, fill="#F8FAFC")
        draw.rounded_rectangle((110, 180, 560, 540), radius=14, outline="#94A3B8", width=3, fill="white")
        draw.rounded_rectangle((640, 180, 1160, 540), radius=14, outline="#94A3B8", width=3, fill="white")
        draw.text((130, 145), "Ô nhập văn bản", fill="#1F4E79", font=font)
        draw.text((660, 145), "Kết quả chuẩn hóa", fill="#1F4E79", font=font)
        draw.text((140, 220), "đk hp mn ơi,\nmk bận rồi T.T", fill="#334E68", font=font)
        draw.text((670, 220), "Đăng ký học phần mọi người ơi,\nmình bận rồi.", fill="#334E68", font=font)
        draw.rounded_rectangle((670, 360, 870, 410), radius=12, outline="#88BDBC", fill="#E0F2F1")
        draw.text((690, 374), "từ đã mở rộng", fill="#2D6A4F", font=small)
    elif kind == "Giao diện quản lý từ điển và kiểm duyệt":
        draw.rounded_rectangle((70, 120, 1210, 615), radius=22, outline="#2F5597", width=4, fill="#F8FAFC")
        draw.text((110, 155), "Từ điển", fill="#1F4E79", font=font)
        draw.text((690, 155), "Kiểm duyệt", fill="#1F4E79", font=font)
        for i, txt in enumerate(["đk: đăng ký", "mn: mọi người", "mk: mình"]):
            y = 210 + i * 75
            draw.rounded_rectangle((110, y, 560, y + 52), radius=10, outline="#CBD5E1", fill="white")
            draw.text((130, y + 14), txt, fill="#334E68", font=small)
        for i, txt in enumerate(["duyệt nghĩa mới", "sửa trước khi lưu", "áp dụng cho lần sau"]):
            y = 210 + i * 75
            draw.rounded_rectangle((690, y, 1140, y + 52), radius=10, outline="#CBD5E1", fill="white")
            draw.text((710, y + 14), txt, fill="#334E68", font=small)

    buffer = BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def add_figure(doc: Document, title: str):
    image = diagram_image(title.replace("Hình 3.1 ", "").replace("Hình 3.2 ", "").replace("Hình 3.3 ", "").replace("Hình 3.4 ", "").replace("Hình 3.5 ", "").replace("Hình 3.6 ", ""))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(image, width=Inches(5.9))
    add_caption(doc, title, above=False)


def add_front_matter(doc: Document, pages: dict[str, int]):
    # Cover page.
    for text, size, bold, after in [
        ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP.HCM", 14, True, 2),
        ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True, 90),
        ("ĐỒ ÁN CƠ SỞ", 18, True, 18),
        ("CÔNG CỤ CHUẨN HÓA VÀ SỬA LỖI TIẾNG VIỆT", 18, True, 54),
        ("Ngành: Công nghệ Thông tin", 13, False, 36),
        ("GVHD: Th.S Nguyễn Thanh Tùng", 13, False, 8),
        ("SVTH: Trương Đông Hồ    MSSV: 2380600754", 13, False, 6),
        ("SVTH: Huỳnh Công Văn    MSSV: 2380606076", 13, False, 6),
        ("SVTH: Võ Ngọc Anh    MSSV: 2380600090", 13, False, 84),
        ("TP. Hồ Chí Minh, tháng 5 năm 2026", 13, False, 0),
    ]:
        p = add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=bold, size=Pt(size), first_line=False, space_after=after)
        p.paragraph_format.line_spacing = 1.2

    doc.add_section(WD_SECTION.NEW_PAGE)
    configure_current_sections(doc)
    body_section = doc.sections[-1]
    body_section.footer.is_linked_to_previous = False
    restart_page_number(body_section, 1)
    add_page_number(body_section)

    add_heading(doc, "LỜI CAM ĐOAN", 1, page_break_before=False)
    for text in [
        'Chúng em cam đoan rằng đồ án "Công cụ chuẩn hóa và sửa lỗi tiếng Việt" là sản phẩm học tập và nghiên cứu do nhóm tự thực hiện dưới sự hướng dẫn của Th.S Nguyễn Thanh Tùng.',
        "Các phần phân tích, thiết kế, cài đặt và đánh giá trong báo cáo được tổng hợp từ quá trình xây dựng hệ thống thực tế của nhóm. Những tài liệu, công cụ và công trình tham khảo được ghi rõ trong phần tài liệu tham khảo.",
        "Nhóm chịu trách nhiệm về nội dung báo cáo, các số liệu trình bày và những kết luận được rút ra từ kết quả thực nghiệm của đồ án.",
    ]:
        add_paragraph(doc, text)
    add_paragraph(doc, "TP. Hồ Chí Minh, tháng 5 năm 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
    add_paragraph(doc, "Sinh viên thực hiện", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)

    doc.add_page_break()
    add_heading(doc, "MỤC LỤC", 1, page_break_before=False)
    toc_table = doc.add_table(rows=0, cols=2)
    toc_table.style = "Table Grid"
    for title, level in TOC_ENTRIES:
        row = toc_table.add_row()
        row.cells[0].text = title
        row.cells[1].text = str(pages.get(title, ""))
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].paragraph_format.left_indent = Cm(0.7 * level)
    style_index_table(toc_table, [5.7, 0.6])
    remove_table_borders(toc_table)

    doc.add_page_break()
    add_heading(doc, "DANH MỤC CÁC KÝ HIỆU, CHỮ VIẾT TẮT", 1, page_break_before=False)
    add_table(
        doc,
        "",
        ["Ký hiệu", "Ý nghĩa sử dụng trong báo cáo"],
        [
            ["AI", "Trí tuệ nhân tạo, dùng để hỗ trợ xử lý các câu khó"],
            ["API", "Cổng giao tiếp để frontend gọi chức năng backend"],
            ["JWT", "Mã xác thực phiên đăng nhập của người dùng"],
            ["NLP", "Xử lý ngôn ngữ tự nhiên"],
            ["REST", "Cách thiết kế API theo tài nguyên và phương thức HTTP"],
            ["SPA", "Ứng dụng web một trang"],
            ["SQL", "Ngôn ngữ truy vấn cơ sở dữ liệu"],
            ["TTL", "Thời gian lưu tạm của một kết quả trong bộ nhớ đệm"],
        ],
        [1.5, 4.8],
    )

    add_heading(doc, "DANH MỤC CÁC BẢNG", 1, page_break_before=False)
    add_index_table(doc, TABLE_LIST, pages)

    add_heading(doc, "DANH MỤC CÁC HÌNH VẼ, ĐỒ THỊ", 1, page_break_before=False)
    add_index_table(doc, FIGURE_LIST, pages)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_index_table(doc: Document, entries: list[str], pages: dict[str, int]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for entry in entries:
        row = table.add_row()
        row.cells[0].text = entry
        row.cells[1].text = str(pages.get(entry, ""))
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_index_table(table, [5.7, 0.6])
    remove_table_borders(table)


def configure_current_sections(doc: Document):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1.27)
        section.footer_distance = Cm(1.27)


def add_chapter_1(doc: Document):
    add_heading(doc, "Chương 1. TỔNG QUAN", 1)
    add_heading(doc, "1.1 Nhiệm vụ đồ án", 2)
    add_heading(doc, "1.1.1 Tính cấp thiết và lý do hình thành đề tài", 3)
    for text in [
        "Trong giao tiếp hằng ngày, đặc biệt trên mạng xã hội và ứng dụng nhắn tin, người Việt thường viết nhanh hơn nhiều so với văn bản học thuật. Một câu có thể bị bỏ dấu, viết tắt, dùng teencode, xen lẫn biểu tượng cảm xúc hoặc kéo dài ký tự để thể hiện cảm xúc. Cách viết này giúp trò chuyện tự nhiên, nhưng lại gây khó khăn khi cần lưu trữ, tìm kiếm, phân tích hoặc dùng lại nội dung.",
        'Ví dụ, câu "đk hp mn ơi, mk bận rồi" vẫn dễ hiểu với nhiều sinh viên, nhưng một hệ thống máy tính sẽ khó xác định chính xác "đk" là "đăng ký", "hp" là "học phần", "mn" là "mọi người" và "mk" là "mình" nếu không có từ điển hoặc ngữ cảnh hỗ trợ. Tương tự, câu không dấu như "mon do an co so nay kho that day" có thể bị hiểu sai nếu chỉ tra từng từ rời rạc.',
        "Từ nhu cầu đó, nhóm chọn đề tài xây dựng công cụ chuẩn hóa và sửa lỗi tiếng Việt. Mục tiêu của đề tài không phải biến mọi câu chat thành văn phong quá trang trọng, mà là giữ cách diễn đạt tự nhiên trong khi khôi phục những phần cần thiết để câu rõ nghĩa hơn. Đây là hướng phù hợp với đồ án cơ sở vì vừa có bài toán xử lý ngôn ngữ, vừa có sản phẩm web có thể chạy thử và đánh giá.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.1.2 Ý nghĩa khoa học và thực tiễn", 3)
    for text in [
        "Về mặt học thuật, đề tài giúp nhóm tiếp cận bài toán xử lý ngôn ngữ tự nhiên tiếng Việt ở mức cơ bản nhưng có tính thực tế cao. Các vấn đề như khôi phục dấu, mở rộng từ viết tắt, xử lý từ mới và lựa chọn nghĩa theo ngữ cảnh đều xuất hiện trực tiếp trong quá trình xây dựng hệ thống.",
        "Về mặt thực tiễn, công cụ có thể hỗ trợ người dùng chuẩn hóa nội dung chat, bình luận hoặc ghi chú nhanh trước khi dùng cho mục đích học tập và làm việc. Hệ thống cũng có phần quản lý từ điển để người dùng hoặc quản trị viên bổ sung nghĩa mới, nhờ đó dữ liệu có thể được cập nhật dần thay vì phụ thuộc hoàn toàn vào một bộ từ điển cố định.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.1.3 Mục tiêu nghiên cứu", 3)
    for text in [
        "Mục tiêu thứ nhất là xây dựng một pipeline chuẩn hóa văn bản tiếng Việt có thể xử lý các lỗi phổ biến trong văn bản chat. Pipeline cần khôi phục dấu tiếng Việt, loại bỏ biểu tượng cảm xúc khi cần, rút gọn ký tự bị kéo dài, mở rộng từ viết tắt và giữ lại cấu trúc câu dễ đọc.",
        "Mục tiêu thứ hai là phát triển một ứng dụng web để người dùng nhập văn bản và nhận kết quả chuẩn hóa trong thời gian gần thực. Giao diện cần đơn giản, có ô nhập, vùng hiển thị kết quả, chức năng sao chép và phần đánh dấu những từ đã được hệ thống mở rộng.",
        "Mục tiêu thứ ba là xây dựng phần backend có API rõ ràng, có xác thực người dùng, lưu lịch sử sử dụng và hỗ trợ quản lý từ điển. Với những trường hợp khó, hệ thống có thể gọi AI để kiểm tra ngữ nghĩa, nhưng vẫn phải có kết quả dự phòng nếu AI bị tắt, hết quota hoặc gặp lỗi mạng.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.1.4 Đối tượng và phạm vi", 3)
    for text in [
        "Đối tượng nghiên cứu của đề tài là văn bản tiếng Việt không chính thức, chủ yếu gồm tin nhắn, bình luận mạng xã hội, câu chat của sinh viên và các câu gõ nhanh thiếu dấu. Đề tài tập trung vào tiếng Việt, không đi sâu vào dịch máy, nhận dạng giọng nói hoặc phân tích cảm xúc.",
        "Phạm vi triển khai của hệ thống gồm frontend React, backend FastAPI, cơ sở dữ liệu SQL Server, bộ dữ liệu từ điển viết tắt và tài nguyên khôi phục dấu. Một số phần nâng cao như huấn luyện mô hình ngôn ngữ riêng, triển khai production quy mô lớn hoặc xử lý văn bản đa ngôn ngữ phức tạp được xem là hướng phát triển sau đồ án.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.2 Cấu trúc đồ án", 2)
    for text in [
        "Báo cáo được trình bày theo bốn chương chính. Chương 1 giới thiệu lý do chọn đề tài, mục tiêu, phạm vi và các hướng nghiên cứu liên quan. Chương này giúp người đọc hiểu vì sao bài toán chuẩn hóa văn bản chat tiếng Việt có ý nghĩa và nhóm giải quyết vấn đề ở mức nào.",
        "Chương 2 trình bày cơ sở lý thuyết và các phương pháp được dùng trong hệ thống. Nội dung gồm đặc điểm tiếng Việt trong văn bản chat, cách chuẩn hóa văn bản, phương pháp khôi phục dấu, xử lý đa nghĩa, kiến trúc phần mềm và những công nghệ chính.",
        "Chương 3 mô tả kết quả thực nghiệm và quá trình cài đặt. Phần này trình bày thiết kế hệ thống, cơ sở dữ liệu, API, pipeline chuẩn hóa, giao diện web, quy trình xây dựng dữ liệu và kết quả kiểm thử.",
        "Chương 4 nêu kết luận, các kết quả đạt được, những hạn chế còn tồn tại và kiến nghị hướng phát triển tiếp theo. Sau bốn chương là tài liệu tham khảo và phụ lục hỗ trợ.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.3 Các nghiên cứu liên quan", 2)
    add_heading(doc, "1.3.1 Khôi phục dấu tiếng Việt", 3)
    for text in [
        "Khôi phục dấu tiếng Việt là bài toán đã được nghiên cứu trong nhiều năm vì tiếng Việt có hệ thống thanh điệu phong phú. Khi người dùng bỏ dấu, một chuỗi ký tự có thể tương ứng với nhiều từ khác nhau. Ví dụ, chuỗi không dấu \"ma\" có thể là \"ma\", \"má\", \"mà\", \"mã\", \"mả\" hoặc \"mạ\" tùy ngữ cảnh.",
        "Các nghiên cứu như ViDiacritics của Nguyen, Nguyen và Dras năm 2022 cung cấp bộ dữ liệu phục vụ đánh giá khôi phục dấu. PhoBERT của Nguyen và Nguyen năm 2020 cũng cho thấy mô hình học sâu có thể đạt kết quả tốt cho nhiều tác vụ tiếng Việt. Tuy nhiên, với đồ án cơ sở, nhóm chọn cách tiếp cận nhẹ hơn để dễ cài đặt, kiểm thử và giải thích.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.3.2 Chuẩn hóa teencode và viết tắt", 3)
    for text in [
        "Teencode và viết tắt trong tiếng Việt thay đổi nhanh theo nhóm người dùng, lứa tuổi và môi trường giao tiếp. Một từ viết tắt có thể mang nhiều nghĩa, còn một nghĩa có thể có nhiều cách viết khác nhau. Vì vậy, cách dùng từ điển thủ công tuy đơn giản nhưng vẫn cần cơ chế cập nhật và kiểm duyệt.",
        "Một số hướng nghiên cứu về chuẩn hóa văn bản mạng xã hội sử dụng học máy để phát hiện từ bất thường, sau đó dùng ngữ cảnh để chọn cách sửa phù hợp. Trong phạm vi đồ án, nhóm kết hợp từ điển viết tắt, cụm từ ưu tiên và phần đề xuất từ mới để hệ thống có khả năng cải thiện dần qua quá trình sử dụng.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.3.3 Ứng dụng AI cho chuẩn hóa văn bản", 3)
    for text in [
        "Các mô hình ngôn ngữ lớn có khả năng hiểu ngữ cảnh tốt, đặc biệt với những câu có nhiều cách hiểu. Tuy vậy, việc gọi AI cho từng ký tự người dùng gõ sẽ làm tăng độ trễ, chi phí và rủi ro phụ thuộc mạng. Đối với một công cụ cần phản hồi nhanh, AI nên được dùng có chọn lọc.",
        "Vì lý do đó, đề tài dùng hướng kết hợp. Những trường hợp rõ ràng được xử lý bằng quy tắc, từ điển và thống kê đơn giản. AI chỉ được gọi khi câu đủ khó hoặc hệ thống thiếu tự tin. Cách làm này phù hợp với mục tiêu đồ án cơ sở vì giúp nhóm cân bằng giữa tốc độ, độ chính xác và khả năng vận hành.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.4 Yêu cầu đặt ra cho hệ thống", 2)
    add_heading(doc, "1.4.1 Yêu cầu về chức năng", 3)
    for text in [
        "Từ mục tiêu của đề tài, hệ thống cần giải quyết được hai nhóm việc. Nhóm thứ nhất là xử lý văn bản: nhận đầu vào từ người dùng, khôi phục dấu, mở rộng từ viết tắt, loại bỏ những biểu tượng không cần thiết và trả về câu dễ đọc hơn. Nhóm thứ hai là quản lý dữ liệu phục vụ chuẩn hóa: lưu từ điển, ghi nhận đề xuất mới, phân quyền quản trị viên và lưu lịch sử sử dụng.",
        "Ở phía người dùng, chức năng quan trọng nhất là màn hình nhập văn bản và xem kết quả. Người dùng không cần hiểu chi tiết pipeline bên trong, nhưng cần nhìn thấy câu sau khi chuẩn hóa, biết phần nào đã được hệ thống thay đổi và có thể sao chép kết quả để dùng tiếp. Nếu một từ được mở rộng chưa đúng ý, người dùng cần có cách đề xuất nghĩa khác.",
        "Ở phía quản trị viên, hệ thống cần hỗ trợ duyệt các nghĩa mới trước khi đưa vào từ điển chung. Đây là bước cần thiết vì ngôn ngữ chat thay đổi nhanh, nhưng không phải mọi đề xuất đều phù hợp cho toàn bộ người dùng. Việc kiểm duyệt giúp dữ liệu phát triển dần mà vẫn giữ được chất lượng.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.4.2 Yêu cầu về chất lượng", 3)
    for text in [
        "Yêu cầu chất lượng đầu tiên là kết quả phải hợp lý với văn bản chat, không chỉ đúng theo văn bản trang trọng. Ví dụ, câu \"mk bận rồi\" nên được chuẩn hóa thành \"mình bận rồi\", không cần biến thành một câu quá nghiêm túc. Nếu hệ thống sửa quá mạnh, người dùng sẽ cảm thấy câu không còn đúng cách nói ban đầu.",
        "Yêu cầu thứ hai là tốc độ. Vì giao diện hoạt động theo kiểu người dùng gõ đến đâu hệ thống xử lý đến đó, backend cần phản hồi trong khoảng thời gian đủ ngắn. Trong đồ án, nhóm đặt mốc tham khảo dưới 200 ms cho luồng từ giao diện đến kết quả. Mốc này không phải chuẩn tuyệt đối, nhưng đủ để đánh giá cảm giác sử dụng khi chạy thử trên máy phát triển.",
        "Yêu cầu thứ ba là khả năng kiểm thử. Mỗi bước xử lý nên có test riêng để khi thay đổi từ điển, sửa pipeline hoặc thêm API mới, nhóm có thể phát hiện lỗi sớm. Vì vậy, báo cáo không chỉ trình bày giao diện hoàn thành mà còn trình bày cách nhóm đo độ chính xác, đo độ trễ và kiểm tra coverage.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "1.4.3 Cách tiếp cận của nhóm", 3)
    for text in [
        "Nhóm không chọn cách dùng AI cho toàn bộ quá trình vì cách đó khó kiểm soát chi phí, khó giải thích từng lỗi và không phù hợp với nhập liệu theo thời gian gần thực. Thay vào đó, hệ thống xử lý phần lớn trường hợp bằng từ điển, quy tắc và thống kê đơn giản. AI chỉ tham gia khi câu có dấu hiệu khó hoặc có nhiều nghĩa cần chọn.",
        "Cách tiếp cận này giúp đồ án có phần lõi rõ ràng: dữ liệu đầu vào đi qua từng bước, mỗi bước có nhiệm vụ cụ thể và có thể kiểm tra độc lập. Khi kết quả sai, nhóm có thể nhìn lại từ điển, cụm từ ưu tiên, bảng bigram hoặc ngưỡng gọi AI để tìm nguyên nhân, thay vì chỉ nhận một câu trả lời từ mô hình bên ngoài.",
    ]:
        add_paragraph(doc, text)


def add_chapter_2(doc: Document):
    add_heading(doc, "Chương 2. CƠ SỞ LÝ THUYẾT", 1)
    add_heading(doc, "2.1 Đặc điểm tiếng Việt trong văn bản chat", 2)
    add_heading(doc, "2.1.1 Đặc điểm ngôn ngữ học", 3)
    for text in [
        "Tiếng Việt là ngôn ngữ đơn âm tiết và có thanh điệu. Dấu thanh giúp phân biệt nghĩa của nhiều từ, nên việc bỏ dấu có thể làm câu trở nên mơ hồ. Trong văn bản chính thức, dấu tiếng Việt thường được viết đầy đủ. Trong tin nhắn nhanh, người dùng lại thường bỏ dấu để tiết kiệm thời gian hoặc do thói quen gõ trên điện thoại.",
        "Văn bản chat còn có nhiều hiện tượng khác như viết tắt, dùng từ lóng, trộn tiếng Anh, kéo dài ký tự và dùng biểu tượng cảm xúc. Những hiện tượng này không hẳn là lỗi trong giao tiếp hằng ngày, nhưng chúng làm dữ liệu khó xử lý hơn nếu đưa vào hệ thống tìm kiếm, thống kê hoặc lưu trữ chính thức.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.1.2 Thách thức kỹ thuật", 3)
    for text in [
        "Thách thức đầu tiên là đa nghĩa. Một từ viết tắt như \"ct\" có thể được hiểu là \"công ty\", \"chương trình\" hoặc một nghĩa khác tùy câu. Nếu hệ thống chỉ tra từ điển theo từng từ đơn lẻ, kết quả dễ sai trong những câu ngắn và thiếu ngữ cảnh.",
        "Thách thức thứ hai là dữ liệu thay đổi liên tục. Từ viết tắt mới xuất hiện theo cộng đồng người dùng, còn các cách nói cũ có thể ít được dùng hơn. Do đó, hệ thống cần cơ chế thêm nghĩa, duyệt nghĩa và ưu tiên nghĩa phù hợp mà không phải sửa trực tiếp trong mã nguồn.",
        "Thách thức thứ ba là yêu cầu phản hồi nhanh. Khi người dùng gõ trên giao diện web, kết quả cần hiển thị gần như ngay lập tức. Vì vậy, pipeline phải nhẹ, có bộ nhớ đệm và chỉ gọi các bước tốn thời gian khi thật sự cần thiết.",
        "Một khó khăn khác nằm ở ranh giới giữa sửa lỗi và giữ văn phong. Câu chat thường có chủ ý thân mật, nên hệ thống không nên sửa mọi thứ theo văn viết trang trọng. Trong đồ án này, nhóm xem chuẩn hóa là quá trình làm câu rõ nghĩa hơn, không phải thay đổi hoàn toàn giọng nói của người dùng.",
        "Ngoài ra, dữ liệu tiếng Việt trên mạng xã hội có nhiều trường hợp không ổn định về cách viết. Cùng một ý có thể được viết thành \"đk\", \"dk\", \"đky\" hoặc \"dang ky\". Nếu chỉ dùng một danh sách từ khóa cứng, hệ thống sẽ bỏ sót nhiều biến thể. Vì vậy, phần dữ liệu cần kết hợp từ điển viết tắt, cụm từ ưu tiên và cơ chế người dùng đề xuất nghĩa mới.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.2 Cơ sở xử lý và chuẩn hóa văn bản", 2)
    add_heading(doc, "2.2.1 Tiền xử lý văn bản", 3)
    for text in [
        "Tiền xử lý là bước đưa văn bản về dạng ổn định trước khi áp dụng các quy tắc sâu hơn. Trong hệ thống, bước này gồm chuẩn hóa Unicode, loại bỏ một số biểu tượng cảm xúc, rút gọn ký tự lặp và tách văn bản thành các token để dễ xử lý.",
        "Chuẩn hóa Unicode cần thiết vì cùng một ký tự tiếng Việt có thể được lưu theo nhiều cách khác nhau. Nếu không chuẩn hóa, hai chuỗi nhìn giống nhau trên màn hình vẫn có thể khác nhau trong bộ nhớ, dẫn đến lỗi tra từ điển hoặc so khớp chuỗi.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.2.2 Mở rộng viết tắt và cụm từ", 3)
    for text in [
        "Mở rộng viết tắt là việc thay một token ngắn bằng dạng đầy đủ hơn. Ví dụ, \"mn\" được mở rộng thành \"mọi người\" và \"đk\" được mở rộng thành \"đăng ký\". Với các cụm nhiều từ như \"do an co so\", hệ thống cần ưu tiên nhận diện cả cụm trước khi xử lý từng từ riêng lẻ.",
        "Để làm được điều đó, hệ thống dùng một chỉ mục cụm từ để tìm các đoạn dài nhất có thể thay thế. Cách này tránh trường hợp xử lý rời rạc làm sai nghĩa. Sau khi xử lý cụm từ, hệ thống mới mở rộng các từ viết tắt đơn còn lại.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.3 Phương pháp khôi phục dấu tiếng Việt", 2)
    add_heading(doc, "2.3.1 Bảng ứng viên và thống kê bigram", 3)
    for text in [
        "Bảng ứng viên là tập ánh xạ từ dạng không dấu sang các dạng có dấu có thể có. Ví dụ, một từ không dấu có thể được ánh xạ sang nhiều từ có dấu khác nhau. Hệ thống không chọn ngẫu nhiên mà dựa vào từ đứng trước để chọn dạng hợp lý hơn.",
        "Thống kê bigram được dùng để ước lượng khả năng hai từ đi liền nhau. Nếu một cặp từ xuất hiện nhiều trong dữ liệu, hệ thống ưu tiên lựa chọn đó. Cách làm này đơn giản hơn mô hình học sâu nhưng có ưu điểm là chạy nhanh và dễ kiểm tra kết quả.",
        "Trong dự án, bảng word_map có 7.022 mục và bảng bigram_freq lưu 200.000 cặp từ phổ biến. Hai tài nguyên này được xây dựng từ bộ dữ liệu khôi phục dấu và được lưu dưới dạng JSON để backend nạp nhanh khi khởi động.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.3.2 Cụm từ ưu tiên theo ngữ cảnh", 3)
    for text in [
        "Một số câu chat không thể xử lý tốt nếu chỉ nhìn từng từ. Ví dụ, \"do an co so\" dễ bị hiểu thành \"đồ ăn cơ sở\" nếu hệ thống dựa quá nhiều vào thống kê thông thường. Trong ngữ cảnh sinh viên, cụm đúng phải là \"đồ án cơ sở\".",
        "Vì vậy, hệ thống bổ sung danh sách cụm từ ưu tiên. Khi phát hiện một cụm đặc biệt, pipeline dùng kết quả đã định nghĩa trước. Cách này giúp sửa các trường hợp thực tế mà dữ liệu lớn chưa bao phủ đầy đủ.",
        "Cụm từ ưu tiên không thay thế hoàn toàn bảng bigram. Nó chỉ xử lý những trường hợp nhóm quan sát thấy dễ sai và có nghĩa rõ trong ngữ cảnh. Cách làm này phù hợp với giai đoạn đồ án cơ sở vì nhóm có thể giải thích vì sao một cụm được ưu tiên, đồng thời vẫn giữ pipeline đơn giản.",
        "Khi mở rộng dữ liệu trong tương lai, danh sách cụm từ ưu tiên nên được kiểm tra bằng regression test. Nếu thêm một cụm mới làm sai các câu cũ, test sẽ giúp phát hiện sớm. Nhờ đó, hệ thống có thể tăng độ bao phủ mà không làm mất ổn định những trường hợp đã xử lý tốt.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.4 Phương pháp giải quyết đa nghĩa", 2)
    add_heading(doc, "2.4.1 Chọn nghĩa bằng quy tắc ngữ cảnh", 3)
    for text in [
        "Đa nghĩa xuất hiện khi một từ viết tắt có nhiều cách mở rộng. Hệ thống xử lý trước bằng quy tắc ngữ cảnh. Cụ thể, nó xem các từ xung quanh token cần mở rộng rồi so khớp với nhóm từ khóa liên quan đến từng nghĩa.",
        "Ví dụ, nếu từ viết tắt nằm gần các từ liên quan đến công việc như \"lương\", \"công việc\" hoặc \"doanh nghiệp\", hệ thống có thể ưu tiên nghĩa \"công ty\". Nếu nằm gần các từ liên quan đến tình cảm, hệ thống có thể giữ lại các phương án khác để người dùng chọn.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.4.2 Kiểm tra bằng AI khi cần", 3)
    for text in [
        "Khi quy tắc không đủ chắc chắn, backend có thể gửi câu và các nghĩa có thể có cho Gemini để chọn phương án phù hợp hơn. Kết quả AI không được xem là nguồn duy nhất. Hệ thống vẫn giữ kết quả dựa trên quy tắc nếu AI lỗi, chậm hoặc vượt giới hạn sử dụng.",
        "Cách dùng AI có chọn lọc giúp giảm chi phí và giữ tốc độ phản hồi. Với đồ án cơ sở, đây cũng là cách tiếp cận dễ trình bày: phần lõi vẫn là pipeline có thể kiểm thử, còn AI chỉ đóng vai trò hỗ trợ trong các trường hợp mơ hồ.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.5 Kiến trúc hệ thống phần mềm", 2)
    add_heading(doc, "2.5.1 Mô hình ba tầng", 3)
    for text in [
        "Hệ thống được tổ chức theo mô hình ba tầng gồm giao diện người dùng, phần xử lý nghiệp vụ và phần dữ liệu. Giao diện dùng React và Vite để người dùng nhập văn bản, xem kết quả và thao tác với từ điển. Backend dùng FastAPI để cung cấp API, xử lý xác thực và chạy pipeline chuẩn hóa. Dữ liệu được lưu trong SQL Server và một số tệp JSON dùng cho khởi tạo.",
        "Cách chia tầng này giúp mỗi phần có nhiệm vụ rõ ràng. Khi cần thay đổi giao diện, nhóm có thể sửa frontend mà không ảnh hưởng trực tiếp đến pipeline. Khi cần mở rộng từ điển hoặc thêm bảng dữ liệu, nhóm có thể cập nhật backend và cơ sở dữ liệu theo từng phần.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.5.2 Dữ liệu, xác thực và lịch sử sử dụng", 3)
    for text in [
        "Phần dữ liệu của hệ thống không chỉ lưu từ điển viết tắt mà còn lưu tài khoản, phiên đăng nhập, lịch sử chuẩn hóa, đề xuất chờ duyệt, cụm từ ưu tiên và thống kê sử dụng. Nhờ đó, hệ thống có thể phục vụ cả người dùng thông thường và quản trị viên.",
        "Xác thực được thực hiện bằng mật khẩu đã băm và token đăng nhập. Khi người dùng đăng xuất, token được ghi nhận để backend có thể thu hồi phiên. Cách làm này an toàn hơn so với việc chỉ để token hết hạn tự nhiên.",
        "Tầng dữ liệu còn hỗ trợ cá nhân hóa ở mức đơn giản. Một người dùng có thể muốn một từ viết tắt mang nghĩa riêng trong nhóm của họ, trong khi quản trị viên chỉ nên đưa những nghĩa phổ biến vào từ điển chung. Việc tách nghĩa cá nhân và nghĩa chung giúp hệ thống linh hoạt hơn mà không làm dữ liệu chung bị nhiễu.",
        "Lịch sử chuẩn hóa được lưu để người dùng xem lại kết quả đã xử lý và để nhóm có cơ sở quan sát các trường hợp thường gặp. Trong phạm vi đồ án, lịch sử này được dùng ở mức chức năng cơ bản. Nếu triển khai rộng hơn, dữ liệu lịch sử cần có thêm chính sách bảo vệ riêng tư và giới hạn thời gian lưu.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.6 Công nghệ sử dụng và ràng buộc triển khai", 2)
    add_heading(doc, "2.6.1 Lý do chọn công nghệ", 3)
    add_table(
        doc,
        "Bảng 2.1 So sánh các phương pháp chuẩn hóa văn bản tiếng Việt",
        ["Phương pháp", "Ưu điểm", "Hạn chế", "Mức phù hợp"],
        [
            ["Từ điển thủ công", "Dễ hiểu, chạy nhanh", "Khó bao phủ hết từ mới", "Phù hợp làm nền"],
            ["Thống kê bigram", "Chọn dấu theo ngữ cảnh gần", "Phụ thuộc dữ liệu huấn luyện", "Phù hợp với đồ án"],
            ["Mô hình học sâu", "Có thể đạt độ chính xác cao", "Cần dữ liệu và tài nguyên lớn", "Để phát triển sau"],
            ["AI gọi qua API", "Hiểu câu mơ hồ tốt hơn", "Có độ trễ và chi phí", "Dùng có chọn lọc"],
            ["Kết hợp quy tắc và AI", "Cân bằng tốc độ, chi phí và độ chính xác", "Cần thiết kế ngưỡng gọi AI", "Phù hợp nhất"],
        ],
        [1.45, 1.75, 1.85, 1.25],
    )
    for text in [
        "Nhóm chọn Python và FastAPI cho backend vì dễ viết API, có hệ sinh thái tốt cho xử lý dữ liệu và phù hợp với kiểm thử tự động. Frontend dùng React vì thuận tiện xây dựng giao diện nhập liệu theo thời gian gần thực. SQL Server được chọn để phù hợp môi trường học tập và dễ quản lý dữ liệu quan hệ.",
        "Đối với phần khôi phục dấu và mở rộng viết tắt, nhóm ưu tiên các phương pháp có thể giải thích được. Điều này giúp việc báo cáo, kiểm thử và sửa lỗi dễ hơn so với chỉ dùng một mô hình AI như hộp đen.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.6.2 Ràng buộc triển khai", 3)
    for text in [
        "Ràng buộc quan trọng nhất là tốc độ phản hồi khi người dùng gõ. Nếu mỗi lần nhập đều gọi AI hoặc truy vấn dữ liệu nặng, giao diện sẽ chậm và khó sử dụng. Vì vậy, backend dùng bộ nhớ đệm ngắn hạn và chỉ gọi những bước tốn thời gian khi có tín hiệu cần thiết.",
        "Ràng buộc thứ hai là tính an toàn dữ liệu. Hệ thống có đăng ký, đăng nhập, phân quyền và lịch sử sử dụng nên cần băm mật khẩu, không lưu token thô và kiểm soát quyền của các API quản trị. Những phần này không phải trọng tâm xử lý ngôn ngữ, nhưng cần đủ chắc để sản phẩm có thể chạy thử nghiêm túc.",
        "Ràng buộc thứ ba là khả năng chạy được trên máy phát triển thông thường. Nhóm không giả định máy luôn có GPU hoặc cơ sở dữ liệu cloud. Vì vậy, tài nguyên khôi phục dấu được lưu thành JSON, cơ sở dữ liệu có thể chạy bằng SQL Server Express hoặc LocalDB, còn AI có thể tắt mà pipeline vẫn trả kết quả.",
        "Ràng buộc cuối cùng là khả năng giải thích. Khi báo cáo một kết quả đúng hoặc sai, nhóm cần chỉ ra được bước nào tạo ra kết quả đó. Đây là lý do hệ thống ưu tiên quy tắc, từ điển và thống kê dễ đọc trước khi dùng AI.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.7 Cơ sở đánh giá kết quả", 2)
    add_heading(doc, "2.7.1 Đánh giá kết quả khôi phục dấu", 3)
    for text in [
        "Để tránh nêu kết quả theo cảm tính, nhóm dùng các chỉ số đo được từ script đánh giá. Với khôi phục dấu, mỗi mẫu kiểm tra gồm một câu không dấu và một câu chuẩn có dấu. Hệ thống khôi phục dấu cho câu không dấu, sau đó kết quả được so sánh với câu chuẩn.",
        "Chỉ số token accuracy được tính theo từng câu trước, sau đó lấy trung bình trên toàn bộ tập mẫu. Công thức sử dụng trong script là: token accuracy bằng tổng của từng tỷ lệ token đúng trong mỗi câu chia cho số câu kiểm tra, sau đó nhân 100%. Với một câu, tỷ lệ token đúng bằng số token khớp với câu chuẩn chia cho tổng số token của câu chuẩn.",
        "Chỉ số sentence exact match nghiêm ngặt hơn. Một câu chỉ được tính đúng nếu toàn bộ chuỗi sau khôi phục giống câu chuẩn sau khi loại bỏ khoảng trắng đầu cuối. Công thức là: sentence exact match bằng số câu khớp hoàn toàn chia cho tổng số câu kiểm tra, sau đó nhân 100%. Vì chỉ cần sai một token là cả câu không khớp, chỉ số này thường thấp hơn token accuracy.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2.7.2 Đánh giá tốc độ và kiểm thử phần mềm", 3)
    for text in [
        "Độ trễ xử lý được đo bằng thời gian bắt đầu và kết thúc của hàm khôi phục dấu trong script đánh giá. Sau khi đo từng mẫu, danh sách độ trễ được sắp xếp tăng dần. Giá trị p50 là độ trễ nằm giữa danh sách, còn p95 là độ trễ tại vị trí 95% của danh sách. Cách đo này giúp nhìn được cả trường hợp thông thường và nhóm trường hợp chậm hơn.",
        "Coverage backend được lấy từ coverage.py. Công thức tổng quát là số dòng hoặc nhánh mã đã được test chạy qua chia cho tổng số dòng hoặc nhánh có thể đo được, sau đó nhân 100%. Trong dự án, `pyproject.toml` đặt `fail_under = 80`, nghĩa là nếu tổng coverage thấp hơn 80% thì lệnh kiểm thử coverage bị xem là không đạt.",
    ]:
        add_paragraph(doc, text)

    add_table(
        doc,
        "Bảng 2.2 Các chỉ số đánh giá được dùng trong đồ án",
        ["Chỉ số", "Cách tính", "Ý nghĩa khi đọc kết quả"],
        [
            ["Token accuracy", "Trung bình của tỷ lệ token đúng trên từng câu", "Cho biết mức độ khôi phục đúng ở cấp từ"],
            ["Sentence exact match", "Số câu khớp hoàn toàn chia cho tổng số câu", "Đánh giá nghiêm ngặt ở cấp câu"],
            ["Latency p50", "Giá trị giữa của danh sách độ trễ đã sắp xếp", "Thể hiện tốc độ trong trường hợp thường gặp"],
            ["Latency p95", "Giá trị tại mốc 95% của danh sách độ trễ", "Cho biết nhóm yêu cầu chậm hơn có còn chấp nhận được không"],
            ["Coverage", "Phần mã được test chạy qua chia cho phần mã có thể đo", "Đánh giá mức độ bao phủ của kiểm thử tự động"],
        ],
        [1.55, 2.55, 2.2],
    )


def add_chapter_3(doc: Document):
    add_heading(doc, "Chương 3. KẾT QUẢ THỰC NGHIỆM", 1)
    add_heading(doc, "3.1 Thiết kế hệ thống", 2)
    add_heading(doc, "3.1.1 Kiến trúc tổng thể", 3)
    for text in [
        "Hệ thống gồm ba phần chính: giao diện web, backend API và cơ sở dữ liệu. Người dùng thao tác trên trình duyệt, frontend gửi yêu cầu đến backend, backend xử lý chuẩn hóa và trả kết quả về giao diện. Khi cần dữ liệu từ điển hoặc lịch sử, backend truy vấn SQL Server. Khi câu đầu vào có dấu hiệu mơ hồ, backend có thể gọi Gemini để hỗ trợ kiểm tra nghĩa.",
        "Điểm quan trọng của kiến trúc là phần chuẩn hóa không bị đặt trực tiếp trong giao diện. Frontend chỉ chịu trách nhiệm nhận văn bản, gửi yêu cầu, hiển thị kết quả và đánh dấu phần được mở rộng. Các quyết định như chọn dấu, mở rộng viết tắt, đọc dữ liệu người dùng hay gọi AI đều nằm ở backend. Cách chia này giúp cùng một API có thể phục vụ thêm CLI hoặc client khác nếu sau này cần mở rộng.",
        "Trong quá trình thiết kế, nhóm ưu tiên luồng dữ liệu dễ lần theo. Một câu nhập vào sẽ đi qua frontend, API, service chuẩn hóa, dữ liệu từ điển và trả về kết quả có cấu trúc. Khi kết quả chưa đúng, nhóm có thể kiểm tra từng điểm trong luồng thay vì chỉ nhìn giao diện cuối cùng. Đây là cách làm phù hợp với đồ án cơ sở vì giúp việc giải thích và sửa lỗi rõ ràng hơn.",
    ]:
        add_paragraph(doc, text)
    add_figure(doc, "Hình 3.1 Kiến trúc tổng thể hệ thống")
    add_table(
        doc,
        "Bảng 3.1 Công nghệ và thư viện sử dụng trong dự án",
        ["Thành phần", "Công nghệ", "Vai trò"],
        [
            ["Backend", "Python, FastAPI, SQLAlchemy, Pydantic", "Xây dựng API, xử lý nghiệp vụ, kiểm tra dữ liệu"],
            ["Frontend", "React, TypeScript, Vite, React Router", "Tạo giao diện web và luồng nhập liệu"],
            ["Cơ sở dữ liệu", "SQL Server hoặc LocalDB", "Lưu tài khoản, từ điển, lịch sử và đề xuất"],
            ["AI", "Gemini API", "Hỗ trợ chọn nghĩa trong câu khó"],
            ["Kiểm thử", "pytest, coverage.py, Vitest, Playwright", "Đảm bảo chức năng chính hoạt động ổn định"],
        ],
        [1.5, 2.35, 2.45],
    )

    add_heading(doc, "3.1.2 Thiết kế cơ sở dữ liệu", 3)
    for text in [
        "Cơ sở dữ liệu được thiết kế theo các nhóm chức năng rõ ràng. Nhóm tài khoản quản lý người dùng và quyền. Nhóm từ điển lưu các từ viết tắt đã được duyệt, nghĩa riêng của từng người dùng và các đề xuất chờ kiểm duyệt. Nhóm lịch sử và thống kê lưu lại kết quả chuẩn hóa để người dùng tra cứu và để quản trị viên theo dõi hệ thống.",
        "Script khởi tạo cơ sở dữ liệu được viết theo hướng có thể chạy lại nhiều lần. Nếu bảng đã tồn tại, script không xóa dữ liệu cũ mà chỉ bổ sung phần còn thiếu. Cách làm này phù hợp với môi trường phát triển vì nhóm thường phải cài lại hoặc cập nhật schema trong quá trình thử nghiệm.",
    ]:
        add_paragraph(doc, text)
    add_figure(doc, "Hình 3.4 Sơ đồ quan hệ dữ liệu rút gọn")
    add_table(
        doc,
        "Bảng 3.2 Các bảng cơ sở dữ liệu chính trong hệ thống",
        ["Tên bảng", "Nhóm chức năng", "Mô tả ngắn"],
        [
            ["dbo.users", "Người dùng", "Lưu tài khoản, mật khẩu đã băm và trạng thái hoạt động"],
            ["dbo.roles, dbo.user_roles", "Phân quyền", "Quản lý quyền user và admin"],
            ["dbo.user_sessions", "Phiên đăng nhập", "Lưu token đã băm để hỗ trợ đăng xuất và thu hồi phiên"],
            ["dbo.abbreviations", "Từ điển", "Lưu từ viết tắt đã được duyệt"],
            ["dbo.pending_abbreviations", "Kiểm duyệt", "Lưu nghĩa mới chờ quản trị viên xử lý"],
            ["dbo.user_abbreviation_overrides", "Cá nhân hóa", "Lưu nghĩa riêng theo từng người dùng"],
            ["dbo.normalization_history", "Lịch sử", "Lưu input, output và thời điểm chuẩn hóa"],
            ["dbo.phrase_overrides", "Cụm từ", "Ưu tiên những cụm nhiều từ dễ bị hiểu sai"],
            ["dbo.abbreviation_usage_stats", "Thống kê", "Ghi nhận từ viết tắt được sử dụng nhiều"],
        ],
        [1.8, 1.45, 3.05],
    )

    add_heading(doc, "3.1.3 Thiết kế REST API", 3)
    for text in [
        "Backend cung cấp các API để frontend đăng nhập, gửi văn bản cần chuẩn hóa, quản lý từ điển, duyệt đề xuất và xem lịch sử. Những API liên quan đến dữ liệu cá nhân hoặc quản trị đều yêu cầu token đăng nhập. Các dữ liệu đầu vào được kiểm tra bằng schema để hạn chế lỗi định dạng.",
        "Các API được chia theo nhóm chức năng để người đọc mã nguồn dễ tìm. Nhóm auth xử lý đăng ký, đăng nhập và đăng xuất. Nhóm normalize xử lý văn bản. Nhóm dictionary và pending phục vụ từ điển, nghĩa riêng và hàng chờ kiểm duyệt. Nhóm history lưu lại kết quả người dùng đã xử lý. Cách chia này chưa phải kiến trúc lớn, nhưng đủ rõ để nhóm phát triển theo từng phần.",
        "Khi trả dữ liệu về frontend, backend cố gắng giữ phản hồi nhất quán. Với chức năng chuẩn hóa, phản hồi không chỉ có chuỗi kết quả mà còn có danh sách từ đã mở rộng, vị trí của chúng và một số thông tin phụ để giao diện đánh dấu. Nhờ vậy, frontend không cần tự đoán từ nào đã bị thay đổi.",
    ]:
        add_paragraph(doc, text)
    add_table(
        doc,
        "Bảng 3.3 Các endpoint REST API chính",
        ["Endpoint", "Phương thức", "Chức năng", "Quyền truy cập"],
        [
            ["/api/auth/login", "POST", "Đăng nhập và nhận token", "Công khai"],
            ["/api/auth/register", "POST", "Tạo tài khoản mới", "Công khai"],
            ["/api/auth/logout", "POST", "Đăng xuất và thu hồi phiên", "Đã đăng nhập"],
            ["/api/normalize/live", "POST", "Chuẩn hóa văn bản khi người dùng nhập", "Đã đăng nhập"],
            ["/api/abbreviations", "GET, POST", "Tra cứu hoặc thêm từ viết tắt", "Theo quyền"],
            ["/api/pending", "GET", "Xem danh sách đề xuất chờ duyệt", "Admin"],
            ["/api/pending/{id}/approve", "POST", "Duyệt đề xuất và lưu vào từ điển", "Admin"],
            ["/api/history", "GET, POST", "Lưu và xem lịch sử chuẩn hóa", "Đã đăng nhập"],
            ["/api/admin/ai-usage", "GET", "Xem thống kê sử dụng AI", "Admin"],
        ],
        [1.65, 0.9, 2.7, 1.05],
    )

    add_heading(doc, "3.1.4 Thiết kế luồng xử lý", 3)
    add_paragraph(doc, "Luồng xử lý chính bắt đầu khi người dùng nhập văn bản. Frontend chờ một khoảng rất ngắn để tránh gửi quá nhiều yêu cầu khi người dùng đang gõ liên tục. Sau đó, frontend gửi văn bản đến backend. Backend lấy dữ liệu từ bộ nhớ đệm hoặc cơ sở dữ liệu, chạy pipeline chuẩn hóa và trả kết quả gồm văn bản sau xử lý cùng danh sách những từ đã được mở rộng.")
    add_paragraph(doc, "Trong luồng live normalize, bộ nhớ đệm có vai trò giảm số lần đọc dữ liệu lặp lại. Dữ liệu từ JSON, SQL Server, nghĩa riêng của người dùng, phrase override và mined phrase được hợp nhất thành một gói dữ liệu dùng cho xử lý. Khi quản trị viên hoặc người dùng thêm nghĩa mới, cache generation được tăng để kết quả cũ không tiếp tục được dùng sai.")
    add_paragraph(doc, "Một điểm khác của luồng xử lý là hệ thống không để lỗi phụ làm hỏng yêu cầu chính. Ví dụ, nếu ghi thống kê sử dụng thất bại, backend vẫn trả kết quả chuẩn hóa. Nếu AI lỗi hoặc quá chậm, hệ thống giữ kết quả rule based. Cách xử lý này giúp người dùng vẫn dùng được chức năng chính ngay cả khi một thành phần phụ gặp vấn đề.")
    add_figure(doc, "Hình 3.3 Luồng xử lý yêu cầu chuẩn hóa theo thời gian gần thực")

    add_heading(doc, "3.1.5 Thiết kế an toàn và xử lý lỗi", 3)
    for text in [
        "Vì hệ thống có đăng nhập và lưu lịch sử, nhóm không xem phần bảo mật là chi tiết phụ. Mật khẩu được băm trước khi lưu, token đăng nhập được ghi nhận theo phiên và khi đăng xuất thì phiên được thu hồi ở phía server. Điều này giúp backend có thể kiểm tra token còn hợp lệ hay không, thay vì chỉ dựa vào thời gian hết hạn.",
        "Các API quản trị như duyệt pending, xem thống kê AI hoặc quản lý dữ liệu chung yêu cầu quyền admin. Người dùng thông thường chỉ thao tác với dữ liệu của mình hoặc gửi đề xuất. Việc phân quyền này cần thiết vì từ điển chung ảnh hưởng đến kết quả chuẩn hóa của nhiều người.",
        "Khi lỗi xảy ra, hệ thống cố gắng trả thông báo dễ hiểu cho giao diện và ghi chi tiết ở phía backend. Ví dụ, lỗi cấu hình cơ sở dữ liệu không nên hiện ra cho người dùng dưới dạng chuỗi kỹ thuật dài. Trong môi trường phát triển, log giúp nhóm tìm nguyên nhân, còn trong giao diện người dùng chỉ cần thông báo rằng hệ thống chưa sẵn sàng hoặc yêu cầu cần thử lại.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.2 Cài đặt pipeline chuẩn hóa", 2)
    add_heading(doc, "3.2.1 Các bước xử lý chính", 3)
    for text in [
        "Pipeline chuẩn hóa được chia thành nhiều bước nhỏ. Bước đầu chuẩn hóa Unicode để văn bản ổn định. Bước tiếp theo loại bỏ emoji và các emoticon không cần thiết. Sau đó, hệ thống khôi phục dấu tiếng Việt, nhận diện cụm từ viết tắt nhiều token, mở rộng từ viết tắt đơn, rút gọn ký tự kéo dài và viết hoa đầu câu.",
        "Việc tách thành từng bước giúp nhóm dễ kiểm thử từng phần. Khi một lỗi xuất hiện, nhóm có thể xác định lỗi nằm ở khôi phục dấu, mở rộng viết tắt hay xử lý giao diện. Đây là cách tổ chức phù hợp với đồ án vì giảm độ phức tạp khi phát triển theo nhóm.",
    ]:
        add_paragraph(doc, text)
    add_figure(doc, "Hình 3.2 Pipeline chuẩn hóa văn bản")

    add_heading(doc, "3.2.2 Xử lý từ viết tắt và cụm từ", 3)
    for text in [
        "Đối với cụm từ, hệ thống ưu tiên so khớp cụm dài nhất trước. Ví dụ, cụm \"do an co so\" được xử lý như một đơn vị để ra \"đồ án cơ sở\". Nếu xử lý từng từ, hệ thống có thể tạo kết quả sai vì \"do\" và \"an\" đều có nhiều cách hiểu.",
        "Đối với từ viết tắt đơn, hệ thống đọc từ điển đã duyệt, nghĩa riêng của người dùng và dữ liệu từ JSON ban đầu. Nếu một từ có nhiều nghĩa, backend có thể trả nhiều biến thể để frontend hiển thị cho người dùng lựa chọn.",
        "Thứ tự ưu tiên được thiết kế theo hướng gần với cách người dùng mong đợi. Nghĩa riêng của người dùng được ưu tiên trong phạm vi tài khoản đó. Phrase override do quản trị viên tạo được ưu tiên khi cần sửa các cụm dễ sai. Sau đó hệ thống mới dùng từ điển chung và dữ liệu bootstrap. Cách sắp xếp này giúp kết quả vừa có tính cá nhân, vừa không làm mất vai trò của dữ liệu chung.",
        "Khi phát hiện một từ lạ đáng nghi, hệ thống không tự động xem mọi từ ngắn là viết tắt. Một số từ thông dụng như tên ứng dụng, từ tiếng Anh hoặc từ đã đúng nghĩa không nên bị ép vào hàng chờ. Vì vậy, pending suggestion chỉ được tạo khi có tín hiệu phù hợp, giúp quản trị viên không bị quá tải bởi các đề xuất không cần thiết.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.2.3 Tích hợp AI có kiểm soát", 3)
    for text in [
        "AI không được gọi cho mọi yêu cầu. Hệ thống chỉ dùng AI khi câu đủ dài, có nhiều thay đổi về dấu hoặc có từ viết tắt mơ hồ. Ngoài ra, kết quả AI được lưu tạm trong bộ nhớ đệm theo nội dung prompt để tránh gọi lại cùng một câu.",
        "Khi Gemini bị lỗi mạng, hết quota hoặc phản hồi chậm, backend giữ kết quả dựa trên quy tắc. Nhờ vậy, người dùng vẫn nhận được kết quả thay vì gặp lỗi hệ thống. Thiết kế này giúp sản phẩm ổn định hơn trong môi trường học tập, nơi API bên ngoài có thể không luôn sẵn sàng.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.2.4 Minh họa quá trình xử lý một câu", 3)
    for text in [
        "Có thể minh họa pipeline bằng câu \"dk hp mn oi, mk ban roi\". Đầu tiên, hệ thống đưa văn bản về Unicode ổn định và tách thành các token. Sau đó, phần khôi phục dấu chuyển các từ có khả năng rõ ràng sang dạng có dấu. Tiếp theo, cụm \"dk hp\" được ưu tiên hiểu là \"đăng ký học phần\" vì đây là cụm thường gặp trong ngữ cảnh sinh viên.",
        "Sau khi xử lý cụm, các từ viết tắt còn lại được mở rộng: \"mn\" thành \"mọi người\", \"mk\" thành \"mình\" và \"roi\" thành \"rồi\" nếu bước khôi phục dấu đã chọn đúng. Kết quả cuối cùng được viết hoa đầu câu và trả kèm danh sách những đoạn đã thay đổi để frontend tô nền. Nhờ trả kèm danh sách này, người dùng có thể thấy hệ thống đã can thiệp ở đâu trong câu.",
        "Ví dụ trên cho thấy pipeline không chỉ thay từng chữ một cách máy móc. Nó cần kết hợp nhiều nguồn: dấu tiếng Việt, cụm từ, viết tắt và trình bày kết quả. Khi một bước chưa đủ tốt, nhóm có thể bổ sung dữ liệu cho đúng bước đó thay vì thay toàn bộ hệ thống.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.3 Cài đặt dữ liệu và tài nguyên khôi phục dấu", 2)
    add_heading(doc, "3.3.1 Quy trình xây dựng dữ liệu", 3)
    for text in [
        "Tài nguyên khôi phục dấu được xây dựng từ bộ dữ liệu ViDiacritics. Script xử lý đọc dữ liệu, tách từ, tạo bảng ánh xạ từ không dấu sang ứng viên có dấu, thống kê các cặp từ thường đi liền nhau và xuất ra các tệp JSON để backend sử dụng.",
        "Ngoài script xử lý thông thường, dự án còn có phiên bản streaming để đọc dữ liệu theo từng phần nhỏ. Cách này giúp tránh tốn quá nhiều bộ nhớ khi dữ liệu lớn và cho phép tiếp tục từ checkpoint nếu quá trình xử lý bị gián đoạn.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.3.2 Kết quả xây dựng dữ liệu", 3)
    for text in [
        "Sau quá trình xây dựng, hệ thống có word_map.json gồm 7.022 mục và bigram_freq.json gồm 200.000 cặp từ phổ biến. Bên cạnh đó, nhóm bổ sung context_phrases.json cho các cụm chat dễ bị hiểu sai và teencode.json cho các từ viết tắt thường gặp.",
        "Các tài nguyên này được nạp vào backend khi chạy hệ thống. Khi quản trị viên duyệt nghĩa mới, dữ liệu có thể được đồng bộ ngược về tệp JSON hoặc lưu trong cơ sở dữ liệu để áp dụng cho các yêu cầu tiếp theo.",
        "Số 7.022 mục của word_map không được ước lượng bằng tay mà lấy từ manifest sinh dữ liệu và có thể kiểm tra lại bằng cách đọc JSON rồi đếm số khóa. Nói cách khác, công thức kiểm tra là số mục word_map bằng số khóa trong `word_map.json`. Khi chạy kiểm tra hiện tại, `len(word_map)` cho kết quả 7022.",
        "Tương tự, số 200.000 cặp từ của bigram_freq được kiểm tra bằng số khóa trong `bigram_freq.json`. Script streaming giữ các cặp từ có tần suất phù hợp và xuất ra manifest với `bigram_entries = 200000`. Vì vậy, đây là số lượng phần tử thực tế trong tệp dữ liệu đang được backend sử dụng, không phải con số mô tả chung.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.3.3 Vai trò của dữ liệu người dùng và dữ liệu quản trị", 3)
    for text in [
        "Ngoài các tệp JSON ban đầu, hệ thống còn có dữ liệu phát sinh trong quá trình sử dụng. Khi người dùng thêm nghĩa riêng, nghĩa đó chỉ ảnh hưởng đến tài khoản của họ. Khi quản trị viên duyệt một đề xuất vào từ điển chung, nghĩa đó mới ảnh hưởng đến nhiều người dùng. Cách tách này giúp hệ thống học thêm từ thực tế nhưng vẫn có bước kiểm soát.",
        "Dữ liệu lịch sử chuẩn hóa giúp người dùng xem lại các câu đã xử lý. Ở góc độ phát triển, nó cũng cho nhóm biết dạng câu nào xuất hiện nhiều và phần nào hay bị sai. Tuy nhiên, vì lịch sử có thể chứa nội dung cá nhân, nếu triển khai thực tế thì cần bổ sung chính sách xóa lịch sử, giới hạn thời gian lưu và thông báo rõ cho người dùng.",
        "Dữ liệu thống kê viết tắt được ghi theo hướng best effort. Nghĩa là nếu ghi thống kê thành công, hệ thống có thêm thông tin để phân tích; nếu ghi thất bại, chức năng chuẩn hóa vẫn không bị dừng. Đây là lựa chọn thực dụng vì thống kê là phần hỗ trợ, còn kết quả chuẩn hóa mới là chức năng chính.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.4 Cài đặt giao diện người dùng", 2)
    add_heading(doc, "3.4.1 Workspace, màn hình chính", 3)
    for text in [
        "Workspace là màn hình người dùng sử dụng nhiều nhất. Giao diện gồm hai vùng chính: vùng nhập văn bản và vùng hiển thị kết quả. Khi người dùng gõ, frontend gọi API sau một khoảng chờ ngắn. Khi người dùng dán một đoạn dài, frontend gửi yêu cầu nhanh hơn để kết quả xuất hiện kịp thời.",
        "Kết quả chuẩn hóa có nút sao chép để người dùng dùng lại nội dung. Những từ được mở rộng từ viết tắt được tô nền nhẹ để người dùng biết hệ thống đã thay đổi phần nào của câu.",
        "Thiết kế hai vùng nhập và kết quả giúp người dùng dễ so sánh trước và sau khi xử lý. Đây là điểm quan trọng với công cụ chuẩn hóa vì người dùng thường muốn kiểm tra xem hệ thống có sửa đúng ý mình không. Nếu kết quả chưa phù hợp, người dùng có thể sửa lại đầu vào hoặc đề xuất nghĩa mới.",
        "Giao diện không cố đưa quá nhiều lựa chọn kỹ thuật ra màn hình chính. Người dùng chỉ cần nhập, xem kết quả và sao chép. Các thao tác quản lý từ điển, kiểm duyệt hoặc cấu hình nâng cao được đặt ở các trang riêng để màn hình chính không bị rối.",
    ]:
        add_paragraph(doc, text)
    add_figure(doc, "Hình 3.5 Giao diện Workspace của người dùng")

    add_heading(doc, "3.4.2 Đánh dấu từ được chuẩn hóa", 3)
    for text in [
        "Backend trả về danh sách các đoạn đã được mở rộng kèm vị trí bắt đầu và kết thúc trong chuỗi kết quả. Frontend dùng thông tin này để đánh dấu đúng phần chữ trên màn hình. Cách làm bằng vị trí giúp tránh đánh dấu nhầm nếu cùng một từ xuất hiện nhiều lần trong câu.",
        "Khi người dùng nhấn vào phần được đánh dấu, giao diện có thể hiển thị lựa chọn hoặc cho phép bổ sung nghĩa mới. Với quản trị viên, nghĩa mới có thể được duyệt vào từ điển chung. Với người dùng thông thường, nghĩa mới có thể được lưu riêng hoặc gửi vào hàng chờ.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.4.3 Trang từ điển và kiểm duyệt", 3)
    for text in [
        "Trang từ điển cho phép tra cứu, thêm và cập nhật nghĩa của từ viết tắt. Quản trị viên có quyền quản lý dữ liệu chung, còn người dùng thông thường có thể tạo nghĩa riêng cho tài khoản của mình. Cách phân quyền này giúp tránh việc một người dùng vô tình làm sai từ điển chung.",
        "Trang kiểm duyệt pending giúp quản trị viên xem các đề xuất mới. Trước khi duyệt, quản trị viên có thể sửa lại nghĩa cho rõ hơn. Sau khi duyệt, nghĩa mới được lưu vào cơ sở dữ liệu và có hiệu lực cho các lần chuẩn hóa sau.",
    ]:
        add_paragraph(doc, text)
    add_figure(doc, "Hình 3.6 Giao diện quản lý từ điển và kiểm duyệt")

    add_heading(doc, "3.4.4 Phản hồi lỗi và trạng thái sử dụng", 3)
    for text in [
        "Trong quá trình sử dụng, giao diện cần phân biệt rõ các trạng thái: đang xử lý, đã có kết quả, chưa đăng nhập, lỗi kết nối hoặc backend chưa sẵn sàng. Nếu mọi lỗi đều hiện như nhau, người dùng sẽ khó biết nên đăng nhập lại, thử lại hay kiểm tra server.",
        "Với đồ án hiện tại, các trạng thái quan trọng đã được xử lý ở mức đủ dùng cho demo. Khi backend trả lỗi xác thực, người dùng được đưa về luồng đăng nhập. Khi yêu cầu chuẩn hóa lỗi, giao diện không làm mất nội dung đang gõ. Cách xử lý này giúp trải nghiệm ổn định hơn trong lúc thử nghiệm.",
        "Ở các phiên bản sau, phần trạng thái có thể được hoàn thiện thêm bằng thông báo nhẹ, lịch sử lỗi gần nhất và hướng dẫn ngắn cho người dùng. Tuy nhiên, báo cáo không xem đây là trọng tâm chính vì mục tiêu đồ án vẫn là pipeline chuẩn hóa và hệ thống web chạy được.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.5 Kết quả kiểm thử và đánh giá", 2)
    add_heading(doc, "3.5.1 Môi trường cài đặt và chạy thử", 3)
    for text in [
        "Hệ thống được chạy thử trên môi trường Windows với backend FastAPI, frontend Vite và SQL Server Express hoặc LocalDB. Backend mặc định chạy ở cổng 8000, frontend chạy ở cổng 5173. Cấu hình cơ sở dữ liệu được đặt trong tệp môi trường và có thể dùng Windows Authentication khi chạy với SQL Server Express.",
        "Quá trình chạy thử gồm tạo môi trường Python, cài thư viện, khởi tạo cơ sở dữ liệu, chạy backend, chạy frontend và đăng nhập bằng tài khoản demo. Khi cả hai server hoạt động, người dùng có thể truy cập giao diện web, nhập văn bản chat và xem kết quả chuẩn hóa.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.5.2 Đánh giá khôi phục dấu và tốc độ xử lý", 3)
    for text in [
        "Phần đánh giá khôi phục dấu sử dụng tệp `data/diacritic/eval_results.json`, được sinh từ script `scripts/eval_diacritic.py`. Tập đánh giá có 1000 mẫu lấy từ `ViDiacritics_test.csv` bằng cách chọn mẫu cố định với seed 42, nhờ đó lần chạy sau có thể đối chiếu với lần chạy trước trên cùng một nhóm câu.",
        "Với mỗi mẫu, script nhận câu không dấu, chạy hàm khôi phục dấu dựa trên word_map và bigram, rồi so sánh kết quả với câu chuẩn có dấu. Token accuracy không được tính bằng cảm nhận thủ công mà bằng công thức trong code: với từng câu, số token đúng chia cho số token chuẩn; sau đó lấy trung bình của 1000 tỷ lệ này và nhân 100%. Giá trị 85,19% trong báo cáo là 0,8519 nhân 100%.",
        "Sentence exact match được tính theo cách nghiêm ngặt hơn. Script đếm số câu có toàn bộ chuỗi khớp đúng với câu chuẩn, sau đó chia cho 1000 mẫu. Kết quả 22,7% tương ứng 227 câu khớp hoàn toàn trên 1000 câu. Chỉ số này thấp hơn token accuracy vì một câu dài chỉ cần sai một token đã không được tính là khớp hoàn toàn.",
        "Độ trễ p50 và p95 được tính từ danh sách thời gian xử lý của từng mẫu. Sau khi sắp xếp danh sách độ trễ, p50 lấy giá trị ở giữa, còn p95 lấy giá trị tại mốc 95%. Với 1000 mẫu, p95 phản ánh nhóm 50 mẫu chậm nhất phía cuối danh sách, nên chỉ số này hữu ích hơn việc chỉ nhìn thời gian trung bình.",
    ]:
        add_paragraph(doc, text)

    add_table(
        doc,
        "Bảng 3.4 Kết quả đánh giá khôi phục dấu trên tập kiểm tra",
        ["Chỉ số", "Cách thu được kết quả", "Kết quả", "Nhận xét"],
        [
            ["Token accuracy", "Trung bình tỷ lệ token đúng trên 1000 mẫu", "85,19%", "Vượt mốc 80% nhóm đặt ra"],
            ["Sentence exact match", "Số câu khớp hoàn toàn chia cho 1000", "22,7%", "Tương ứng 227 câu đúng toàn bộ"],
            ["Độ trễ p50", "Giá trị giữa của danh sách 1000 độ trễ", "0,1 ms", "Rất nhanh với xử lý rule based"],
            ["Độ trễ p95", "Giá trị tại mốc 95% sau khi sắp xếp độ trễ", "0,5 ms", "Nhóm yêu cầu chậm vẫn dưới 1 ms"],
            ["Nguồn dữ liệu", "ViDiacritics_test.csv, sample_size = 1000", "seed 42", "Có thể chạy lại để đối chiếu"],
        ],
        [1.35, 2.25, 1.0, 1.7],
    )
    for text in [
        "Kết quả trên cho thấy pipeline rule based đã đạt mức đủ tốt cho mục tiêu đồ án cơ sở. Điểm mạnh của cách làm này là tốc độ rất thấp và dễ kiểm tra. Điểm yếu là những câu có nhiều nghĩa hoặc câu chat quá xa dữ liệu huấn luyện vẫn có thể sai, nên hệ thống cần thêm cụm từ ưu tiên và AI hỗ trợ khi thiếu tự tin.",
        "Khi so sánh token accuracy với sentence exact match, nhóm không xem 22,7% là dấu hiệu hệ thống thất bại. Với một câu gồm nhiều token, xác suất để tất cả token đều đúng luôn thấp hơn xác suất từng token đúng riêng lẻ. Vì vậy, hai chỉ số được đọc cùng nhau: token accuracy cho biết chất lượng ở cấp từ, còn sentence exact match cho biết mức độ hoàn chỉnh ở cấp câu.",
    ]:
        add_paragraph(doc, text)

    add_table(
        doc,
        "Bảng 3.5 Kết quả regression trên tập câu chat thực tế",
        ["Nhóm câu", "Số mẫu", "Token accuracy trung bình", "Số câu khớp hoàn toàn"],
        [
            ["ambiguous", "10", "73,5%", "4/10"],
            ["basic", "10", "87,0%", "5/10"],
            ["context_phrase", "10", "90,8%", "7/10"],
            ["formal", "5", "89,3%", "3/5"],
            ["mixed_english", "5", "100,0%", "5/5"],
            ["natural_chat", "4", "78,6%", "1/4"],
            ["short", "7", "100,0%", "7/7"],
            ["Tổng hợp", "51", "87,7%", "32/51"],
        ],
        [1.55, 0.8, 1.75, 1.75],
    )
    for text in [
        "Tập regression chat gồm 51 câu trong `data/diacritic/chat_regression_test.csv`. Công thức tính token accuracy vẫn giống phần đánh giá chính: mỗi câu được chấm theo tỷ lệ token đúng, sau đó lấy trung bình. Kết quả tổng hợp 87,7% là trung bình của 51 điểm token accuracy theo từng câu.",
        "Số câu khớp hoàn toàn trong tập chat là 32 trên 51, tương đương 62,7% nếu tính theo công thức 32 chia 51 rồi nhân 100%. Nhóm không đưa tỷ lệ này vào bảng chính vì bảng đã ghi rõ 32/51, nhưng khi cần diễn giải, đây là cách tính phần trăm tương ứng.",
        "Các nhóm mixed_english và short đạt 100,0% vì câu ngắn, cấu trúc rõ hoặc ít mơ hồ. Nhóm ambiguous thấp hơn, 73,5%, vì chứa những từ có nhiều cách hiểu. Kết quả này phù hợp với thiết kế của hệ thống: phần rule based xử lý nhanh các câu rõ ràng, còn câu mơ hồ cần AI hoặc người dùng chọn biến thể.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.5.3 Kết quả kiểm thử phần mềm", 3)
    for text in [
        "Bên cạnh đánh giá chất lượng ngôn ngữ, nhóm kiểm thử phần mềm để kiểm tra các luồng chính của backend và frontend. Kết quả test được đọc theo nguyên tắc rõ ràng: pass rate của test bằng số test pass chia cho tổng số test thực thi, còn test skipped được ghi nhận riêng vì nó không phải lỗi thất bại.",
        "Với backend, kết quả `375 passed, 1 skipped, 36 subtests` nghĩa là 375 test chính đã chạy thành công, 1 test được bỏ qua theo điều kiện, và 36 subtests là các trường hợp con trong một số test. Nếu chỉ tính các test đã thực thi, pass rate là 375 chia 375, bằng 100%. Nếu tính cả test skipped vào tổng số test được thu thập, tỷ lệ pass là 375 chia 376, xấp xỉ 99,73%, nhưng skipped không được xem là lỗi trong pytest.",
        "Với frontend, kết quả 29 passed nghĩa là 29 test Vitest đều chạy thành công. Pass rate được tính là 29 chia 29, bằng 100%. Cách trình bày này giúp người đọc hiểu rõ tỷ lệ đạt đến từ phép chia nào, không chỉ là một dòng trạng thái chung chung.",
    ]:
        add_paragraph(doc, text)
    add_table(
        doc,
        "Bảng 3.6 Tổng hợp kết quả kiểm thử phần mềm",
        ["Nhóm kiểm thử", "Kết quả hiện tại", "Công cụ", "Ý nghĩa"],
        [
            ["Backend", "375 passed, 1 skipped, 36 subtests", "pytest", "Kiểm tra route, service, repository và pipeline"],
            ["Coverage backend", "82%", "coverage.py", "Vượt ngưỡng 80% đã đặt trong cấu hình"],
            ["Ruff", "Pass", "ruff", "Kiểm tra quy ước và lỗi tĩnh cơ bản"],
            ["Mypy", "Pass trong phạm vi cấu hình", "mypy", "Kiểm tra kiểu ở các module quan trọng"],
            ["Bandit", "Pass", "bandit", "Kiểm tra bảo mật tĩnh"],
            ["Frontend", "29 passed", "Vitest", "Kiểm tra component và hook chính"],
            ["Frontend build", "Pass", "Vite", "Đảm bảo giao diện có thể build"],
        ],
        [1.45, 1.7, 1.15, 2.0],
    )
    for text in [
        "Bộ kiểm thử giúp nhóm tự tin hơn khi sửa pipeline hoặc thêm chức năng mới. Các test repository dùng fake session để không bắt buộc phải có SQL Server thật trong mọi nhánh kiểm thử. Đây là lựa chọn hợp lý vì môi trường phát triển của từng máy có thể khác nhau.",
        "Coverage backend đạt 82% theo báo cáo của coverage.py. Công thức chung là phần mã đã được test chạy qua chia cho phần mã có thể đo được, sau đó nhân 100%. Ngưỡng trong `pyproject.toml` là 80%, nên 82% cao hơn ngưỡng yêu cầu 2 điểm phần trăm. Điều này không có nghĩa là hệ thống hết lỗi, nhưng cho thấy phần lớn nhánh quan trọng đã có kiểm thử tự động.",
        "Ngoài các test tự động, nhóm cũng kiểm tra thủ công các luồng chính: đăng nhập, nhập câu chat, xem kết quả, sao chép kết quả, thêm nghĩa mới và duyệt pending. Những luồng này là phần người dùng trực tiếp thao tác nên cần được kiểm tra trước khi kết luận hệ thống hoàn thành.",
    ]:
        add_paragraph(doc, text)

    add_table(
        doc,
        "Bảng 3.7 Cách đối chiếu các chỉ số phần mềm",
        ["Nội dung", "Cách tính hoặc cách kiểm tra", "Kết luận"],
        [
            ["Backend pass rate", "375 test pass chia cho 375 test thực thi", "100% trên phần đã chạy"],
            ["Backend pass nếu tính skipped", "375 chia cho 376 test được thu thập", "Xấp xỉ 99,73%, skipped không phải lỗi"],
            ["Frontend pass rate", "29 test pass chia cho 29 test thực thi", "100%"],
            ["Coverage backend", "Dòng mã đã chạy qua test chia cho dòng mã có thể đo", "82%, cao hơn ngưỡng 80%"],
            ["Quality gate", "Ruff, mypy, bandit và build frontend đều pass", "Đạt các cổng kiểm tra hiện tại"],
        ],
        [1.55, 3.0, 1.75],
    )

    add_heading(doc, "3.5.4 Đối chiếu với yêu cầu ban đầu", 3)
    for text in [
        "Sau khi có kết quả thực nghiệm, nhóm đối chiếu lại với các yêu cầu đã nêu ở Chương 1. Việc đối chiếu này giúp báo cáo không chỉ liệt kê chức năng mà còn cho thấy mỗi phần đã phục vụ mục tiêu nào của đồ án.",
        "Nhóm không dùng bảng này để khẳng định hệ thống hoàn hảo. Bảng chỉ cho thấy trạng thái hiện tại: phần nào đã hoàn thành ở mức đồ án cơ sở, phần nào còn cần phát triển tiếp nếu muốn đưa vào sử dụng rộng hơn.",
    ]:
        add_paragraph(doc, text)

    add_table(
        doc,
        "Bảng 3.8 Đối chiếu yêu cầu ban đầu và kết quả thực hiện",
        ["Yêu cầu", "Kết quả đã thực hiện", "Nhận xét"],
        [
            ["Chuẩn hóa văn bản chat", "Có pipeline xử lý dấu, viết tắt, emoji, ký tự kéo dài và cụm từ ưu tiên", "Đạt mục tiêu chính ở phạm vi đồ án"],
            ["Giữ văn phong tự nhiên", "Không ép toàn bộ câu thành văn phong trang trọng", "Phù hợp định hướng humanize của đề tài"],
            ["Giao diện sử dụng được", "Có workspace nhập văn bản, xem kết quả, sao chép và highlight", "Đủ để chạy demo và kiểm tra trực tiếp"],
            ["Quản lý từ điển", "Có dữ liệu viết tắt, nghĩa riêng, pending và quyền admin", "Cần hoàn thiện thêm dashboard nếu triển khai rộng"],
            ["Đánh giá bằng số liệu", "Có token accuracy, sentence exact, latency, regression chat và coverage", "Các tỷ lệ được trình bày kèm cách tính"],
            ["Có khả năng mở rộng", "Dữ liệu tách khỏi mã nguồn, backend chia theo service và repository", "Có nền tảng để phát triển tiếp"],
        ],
        [1.6, 3.0, 1.7],
    )

    add_heading(doc, "3.6 Nhận xét kết quả thực nghiệm", 2)
    add_heading(doc, "3.6.1 Những kết quả phù hợp với mục tiêu ban đầu", 3)
    for text in [
        "Kết quả thực nghiệm cho thấy hệ thống đạt được mục tiêu chính của đồ án: có thể chuẩn hóa câu chat tiếng Việt ở mức sử dụng được, phản hồi nhanh và có giao diện để người dùng thao tác trực tiếp. Phần khôi phục dấu đạt 85,19% trên 1000 mẫu kiểm tra và 87,7% trên tập câu chat regression, tức là kết quả không chỉ đúng trên dữ liệu báo chí mà còn có tín hiệu tốt trên nhóm câu gần với cách người dùng gõ thật.",
        "Về phần mềm, backend và frontend đều có test tự động. Điều này quan trọng vì hệ thống không chỉ là một đoạn script xử lý văn bản. Nó có đăng nhập, phân quyền, cơ sở dữ liệu, API, giao diện và các luồng cập nhật từ điển. Khi nhiều phần nối với nhau, test giúp giảm rủi ro sửa một chỗ nhưng làm hỏng chỗ khác.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "3.6.2 Những điểm cần đọc thận trọng", 3)
    for text in [
        "Một số chỉ số cần được hiểu đúng phạm vi. Token accuracy 85,19% không có nghĩa mọi câu đều đúng hoàn toàn. Nó cho biết trung bình cấp token đạt mức đó. Vì vậy, người dùng vẫn có thể gặp một câu sai ở vài từ, đặc biệt khi câu chứa từ lóng, tên riêng hoặc cách viết quá mới.",
        "Sentence exact match 22,7% cũng không nên bị đọc như một tỷ lệ thất bại đơn giản. Đây là chỉ số rất nghiêm ngặt. Trong câu dài, chỉ một từ sai dấu hoặc sai dấu câu đã làm câu không khớp hoàn toàn. Nhóm giữ chỉ số này trong báo cáo để nhìn thẳng vào hạn chế của hệ thống, không chỉ chọn chỉ số có vẻ đẹp hơn.",
        "Kết quả regression chat cho thấy nhóm ambiguous thấp hơn các nhóm khác. Điều này đúng với bản chất bài toán. Các câu mơ hồ thường không thể giải quyết chắc chắn bằng từ điển và bigram. Hướng xử lý hợp lý là cho AI hoặc người dùng hỗ trợ chọn nghĩa, thay vì ép hệ thống luôn chọn một đáp án duy nhất.",
    ]:
        add_paragraph(doc, text)


def add_chapter_4(doc: Document):
    add_heading(doc, "Chương 4. KẾT LUẬN VÀ KIẾN NGHỊ", 1)
    add_heading(doc, "4.1 Kết luận", 2)
    add_heading(doc, "4.1.1 Kết quả đạt được", 3)
    for text in [
        "Đồ án đã xây dựng được một hệ thống chuẩn hóa và sửa lỗi tiếng Việt có thể chạy thử qua giao diện web. Hệ thống xử lý được các trường hợp phổ biến trong văn bản chat như bỏ dấu, viết tắt, teencode, emoji, ký tự kéo dài và một số cụm từ dễ hiểu sai.",
        "Về mặt phần mềm, nhóm đã hoàn thành các thành phần chính gồm frontend React, backend FastAPI, cơ sở dữ liệu SQL Server, hệ thống đăng nhập, quản lý từ điển, kiểm duyệt pending, lưu lịch sử và thống kê sử dụng. Các API chính đã có kiểm thử và có thể phục vụ giao diện hiện tại.",
        "Về mặt đánh giá, kết quả khôi phục dấu đạt token accuracy 85,19% theo công thức trung bình tỷ lệ token đúng trên 1000 mẫu. Tập regression chat đạt 87,7% theo cùng cách tính trên 51 câu gần với cách gõ thực tế. Bộ kiểm thử backend đạt coverage 82%, cao hơn ngưỡng 80% trong cấu hình dự án.",
        "Những kết quả này cho thấy hệ thống không chỉ dừng ở mức ý tưởng. Nhóm đã có sản phẩm chạy được, có dữ liệu kiểm thử, có công thức đánh giá và có các cổng kiểm tra chất lượng cơ bản. Dù vậy, kết quả vẫn cần được hiểu trong phạm vi đồ án cơ sở: dữ liệu chưa đủ lớn để khẳng định hệ thống xử lý tốt mọi dạng văn bản tiếng Việt ngoài thực tế.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "4.1.2 Những điểm còn hạn chế", 3)
    for text in [
        "Hệ thống vẫn còn một số hạn chế. Thứ nhất, dữ liệu viết tắt và cụm từ chat chưa thể bao phủ hết cách nói của nhiều nhóm người dùng khác nhau. Khi gặp từ mới, hệ thống cần người dùng hoặc quản trị viên bổ sung nghĩa để cải thiện dần.",
        "Thứ hai, phần AI hiện phụ thuộc vào API bên ngoài nên có thể bị ảnh hưởng bởi mạng, quota hoặc thay đổi cấu hình. Dù hệ thống đã có fallback, chất lượng của các câu mơ hồ vẫn có thể giảm khi AI không khả dụng.",
        "Thứ ba, giao diện quản trị cho một số dữ liệu như thống kê sử dụng, phiên đăng nhập và tùy chọn người dùng chưa đầy đủ. Các API đã có nền tảng, nhưng vẫn cần thêm màn hình để quản trị viên theo dõi thuận tiện hơn.",
        "Thứ tư, phần đánh giá hiện chủ yếu dựa trên token accuracy, sentence exact, regression chat và test phần mềm. Các chỉ số này đủ cho đồ án cơ sở, nhưng chưa phản ánh đầy đủ cảm nhận của người dùng. Ví dụ, một câu có thể sai một từ nhưng người đọc vẫn hiểu đúng ý, hoặc ngược lại, một câu đúng nhiều token nhưng sai đúng từ quan trọng. Vì vậy, giai đoạn sau nên bổ sung đánh giá thủ công trên một tập câu thật.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "4.2 Kiến nghị", 2)
    add_heading(doc, "4.2.1 Hướng phát triển kỹ thuật", 3)
    for text in [
        "Trong giai đoạn tiếp theo, hệ thống nên tách route backend theo từng nhóm chức năng như auth, normalize, dictionary, pending, history và admin. Việc tách module giúp mã nguồn dễ đọc hơn, đặc biệt khi số lượng API tăng.",
        "Cơ sở dữ liệu nên chuyển dần từ script khởi tạo tổng hợp sang công cụ migration như Alembic. Cách này giúp theo dõi lịch sử thay đổi schema rõ ràng hơn và giảm rủi ro khi cập nhật trên nhiều máy.",
        "Về bảo mật, nếu hệ thống được đưa ra môi trường public, nhóm nên cân nhắc chuyển token từ sessionStorage sang cookie httpOnly kết hợp SameSite và CSRF phù hợp. Ngoài ra, cần bổ sung quét lỗ hổng thư viện trong CI khi có môi trường mạng ổn định.",
        "Về đánh giá, nhóm nên giữ nguyên các công thức hiện có để tiện so sánh qua từng phiên bản. Khi thêm dữ liệu mới, kết quả cần được báo cáo cùng số mẫu, cách chọn mẫu và thời điểm chạy. Nếu chỉ ghi một tỷ lệ phần trăm mà không ghi cách tính, người đọc sẽ khó biết tỷ lệ đó phản ánh điều gì.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "4.2.2 Hướng phát triển sản phẩm", 3)
    for text in [
        "Về chất lượng chuẩn hóa, hướng quan trọng nhất là mở rộng tập câu chat thực tế để giảm lỗi trong những ngữ cảnh đời thường. Dữ liệu càng gần cách người dùng viết thật thì hệ thống càng có khả năng chọn dấu và nghĩa chính xác hơn.",
        "Về giao diện, hệ thống nên bổ sung dashboard cho quản trị viên, quản lý phiên đăng nhập, tùy chọn cá nhân và thống kê AI. Các luồng quan trọng như đánh dấu từ được mở rộng, popup thêm nghĩa và duyệt pending cũng nên có kiểm thử trình duyệt đầy đủ hơn.",
        "Nếu có thêm thời gian và dữ liệu, nhóm có thể thử huấn luyện hoặc tinh chỉnh mô hình tiếng Việt nhỏ cho bài toán khôi phục dấu và chọn nghĩa. Tuy nhiên, hướng này cần được thực hiện sau khi pipeline hiện tại đã ổn định và có đủ dữ liệu đánh giá đáng tin cậy.",
    ]:
        add_paragraph(doc, text)


def add_references_and_appendix(doc: Document):
    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    references = [
        "FastAPI Documentation. Sebastián Ramírez. Truy cập ngày 01/05/2026. https://fastapi.tiangolo.com.",
        "Google Gemini API Documentation. Google LLC. Truy cập ngày 01/05/2026. https://ai.google.dev/gemini-api.",
        "Jones, M.B., Bradley, J., và Sakimura, N. năm 2015. JSON Web Token. RFC 7519. Internet Engineering Task Force.",
        "Nguyen, D.Q., và Nguyen, A.T. năm 2020. PhoBERT: Pre-trained language models for Vietnamese. Findings of the Association for Computational Linguistics: EMNLP 2020.",
        "Nguyen, D.Q., Nguyen, A.T., và Dras, M. năm 2022. ViDiacritics: A Vietnamese Diacritic Restoration Dataset. Proceedings of the 13th Language Resources and Evaluation Conference.",
        "Provos, N., và Mazières, D. năm 1999. A Future-Adaptable Password Scheme. Proceedings of the 1999 USENIX Annual Technical Conference.",
        "React Documentation. Meta Open Source. Truy cập ngày 01/05/2026. https://react.dev.",
        "SQLAlchemy Documentation. Mike Bayer. Truy cập ngày 01/05/2026. https://docs.sqlalchemy.org.",
        "Tran, T.D., Nguyen, T.L., và Le, H. năm 2019. Lexical Normalization for Vietnamese Social Media Text. Proceedings of Pacific Asia Conference on Language, Information and Computation.",
        "Unicode Emoji Test File. Unicode Consortium. Truy cập ngày 01/05/2026. https://unicode.org/emoji/charts/emoji-list.html.",
    ]
    for item in references:
        add_paragraph(doc, item, first_line=False)

    add_heading(doc, "PHỤ LỤC", 1)
    add_heading(doc, "Phụ lục A. Cấu trúc thư mục dự án", 2)
    appendix_lines = [
        "ChuanHoaTiengViet/",
        "app/: chứa backend FastAPI, xác thực, pipeline chuẩn hóa và repository dữ liệu.",
        "frontend/: chứa giao diện React, component, hook và kiểm thử frontend.",
        "database/: chứa script khởi tạo schema và công cụ seed dữ liệu.",
        "data/: chứa từ điển viết tắt, tài nguyên khôi phục dấu và dữ liệu hỗ trợ.",
        "scripts/: chứa script chạy server, build asset và đánh giá.",
        "tests/: chứa bộ kiểm thử backend bằng pytest.",
    ]
    for line in appendix_lines:
        add_paragraph(doc, line, first_line=False)

    add_heading(doc, "Phụ lục B. Ví dụ chuẩn hóa tiêu biểu", 2)
    add_table(
        doc,
        "Bảng B.1 Ví dụ chuẩn hóa tiêu biểu",
        ["Văn bản đầu vào", "Kết quả sau chuẩn hóa"],
        [
            ["đk hp mn ơi, mk bận rồi T.T", "Đăng ký học phần mọi người ơi, mình bận rồi."],
            ["toi di hoc, gap ct nhe", "Tôi đi học, gặp công ty nhé hoặc một nghĩa khác tùy ngữ cảnh."],
            ["đẹppppp quá trờiiii", "Đẹp quá trời."],
            ["mon do an co so nay kho that day", "Môn đồ án cơ sở này khó thật đấy."],
            ["ntn b ơi, cb xong chua", "Như thế nào bạn ơi, chuẩn bị xong chưa."],
            ["bff mk mới quen thấy ổn ghê", "Bạn thân mình mới quen thấy ổn ghê."],
        ],
        [2.95, 3.35],
    )

    add_heading(doc, "Phụ lục C. Cấu hình môi trường mẫu", 2)
    config = [
        "SQL_MODE=localdb",
        "SQL_INSTANCE=MSSQLLocalDB",
        "SQL_DATABASE=VietNormalizer",
        "AI_PROVIDER=gemini",
        "GEMINI_MODEL=gemini-2.5-flash-lite",
        "AI_TIMEOUT_SECONDS=10",
        "AI_RATE_LIMIT_COOLDOWN=60",
        "SEMANTIC_VERIFY_ENABLED=1",
        "AUTH_SECRET=<chuoi-bi-mat-toi-thieu-32-ky-tu>",
    ]
    for line in config:
        add_paragraph(doc, line, first_line=False)


def sanitize_doc_text(doc: Document):
    forbidden = {
        "—": ",",
        "–": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if not run.text:
                continue
            for old, new in forbidden.items():
                run.text = run.text.replace(old, new)
            run.text = re.sub(r"\[\d+\]", "", run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        for old, new in forbidden.items():
                            run.text = run.text.replace(old, new)
                        run.text = re.sub(r"\[\d+\]", "", run.text)


def build_docx(pages: dict[str, int]):
    doc = Document()
    configure_document(doc)
    add_front_matter(doc, pages)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_references_and_appendix(doc)
    sanitize_doc_text(doc)
    configure_current_sections(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pages-json", default="")
    args = parser.parse_args()
    pages: dict[str, int] = {}
    if args.pages_json:
        pages = json.loads(Path(args.pages_json).read_text(encoding="utf-8"))
    build_docx(pages)


if __name__ == "__main__":
    main()
